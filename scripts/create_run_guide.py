"""Build the handoff run guide using the bundled document runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output/ML-IDS-Run-Guide.docx'
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.68)
section.left_margin = section.right_margin = Inches(.72)
section.footer_distance = Inches(.3)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number'):
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(10.5)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.08
for name, size in [('Title', 26), ('Heading 1', 19), ('Heading 2', 13), ('Heading 3', 11)]:
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.bold = True
    doc.styles[name].paragraph_format.space_before = Pt(10 if name != 'Title' else 0)
    doc.styles[name].paragraph_format.space_after = Pt(7)
    doc.styles[name].paragraph_format.keep_with_next = True
doc.styles['Subtitle'].font.size = Pt(12)
for border in list(doc.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
doc.core_properties.title = 'ML IDS Setup and Run Guide'
doc.core_properties.subject = 'Running the PyTorch and Streamlit packet intrusion detection project'
doc.core_properties.author = 'ML IDS Project'
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run('ML IDS Run Guide  |  18 September 2026  |  ')
run.font.size = Pt(8)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)


def p(text, bold_lead=None):
    para = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        para.add_run(bold_lead).bold = True
        para.add_run(text[len(bold_lead):])
    else:
        para.add_run(text)
    return para


def h(text, level=2):
    return doc.add_heading(text, level=level)


def code(text):
    para = doc.add_paragraph()
    para.paragraph_format.left_indent = Inches(.12)
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.line_spacing = 1.05
    para.paragraph_format.keep_together = True
    run = para.add_run(text)
    run.font.name = 'Liberation Mono'
    run.font.size = Pt(8.5)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Liberation Mono')
    return para


def bullet(text):
    return doc.add_paragraph(text, style='List Bullet')


def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
    for row in rows:
        for cell, text in zip(t.add_row().cells, row):
            cell.text = text
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        for key, value in [('val','single'), ('sz','5'), ('color','D9D9D9')]:
            el.set(qn('w:' + key), value)
        borders.append(el)
    t._tbl.tblPr.append(borders)
    for i, row in enumerate(t.rows):
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement('w:cantSplit'))
        if i == 0:
            props.append(OxmlElement('w:tblHeader'))
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                el = OxmlElement('w:' + edge); el.set(qn('w:w'), '100'); el.set(qn('w:type'), 'dxa'); margins.append(el)
            tcpr.append(margins)
            if i == 0:
                shading = OxmlElement('w:shd'); shading.set(qn('w:fill'),'E8EFF8'); tcpr.append(shading)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs:
                    run.font.size = Pt(9.5)
                    run.bold = i == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def page(title):
    doc.add_page_break()
    h(title, 1)


doc.add_paragraph('ML IDS Setup and Run Guide', 'Title')
doc.add_paragraph('PyTorch packet classification and the local Streamlit dashboard', 'Subtitle')
p('Use this guide to start the application, inspect packet bytes, prepare CICIoT2023 captures and train the CNN BiLSTM model. The quickest route on your current Mac is below. New installations start on page 2; data preparation and training start on page 3.')
p('Current status: the application is implemented and its 26 software checks passed. The supplied CICIOT23 folder contains three flow CSVs with 46 features and no PCAP files. Real packet-model training and detection remain pending labeled raw captures. The supplied package contains no trained packet checkpoint.', 'Current status:')
h('1 Open Terminal and enter the project folder')
code('cd /Users/harsha/Downloads/ML-IDS')
p('Use this original project folder on your Mac. Its Python environment is already installed in .venv-pytorch. Do not recreate that environment just to launch the dashboard.')
h('2 Start the dashboard')
code('./start.sh')
p('Keep this Terminal window open. You should see a Streamlit message with the local address. You can also double-click Start-IDS.command in Finder. If the dashboard is already running, open its address without launching another copy.')
h('3 Open the local address')
code('http://127.0.0.1:8501')
p('The current app has a blue light theme, 30 px action buttons, rounded cards and no shadows. The sidebar should show “Packet model awaits training” until real training is completed.')
h('4 Try the available packet inspector')
p('Choose Packet inspector, leave Hex payload selected, enter the following bytes and click Inspect payload:')
code('00 7f ff')
p('Expected result: 3 original bytes, 3 retained bytes and 1,021 zero-padding bytes, forming a 1,024-value signal. The observed values are 0, about 0.498 and 1. Classification and SHAP remain disabled until a compatible trained model is loaded.')
h('5 Stop or restart the application')
p('Press Control+C in the Terminal window running Streamlit to stop it. Run ./start.sh again to restart. Closing only the browser tab does not stop the server.')

page('Install a fresh copy')
p('Use these steps on a different computer or after extracting ML-IDS-PyTorch-Streamlit.zip. The archive contains source code and instructions; Python packages and the raw dataset are installed separately. Python 3.12 is the verified version.')
h('On macOS')
p('1. Extract the ZIP. Open Terminal and change into its ML-IDS-PyTorch-Streamlit folder. If it is in Downloads, use:')
code('cd ~/Downloads/ML-IDS-PyTorch-Streamlit')
p('2. Create the environment and install the dependencies. Internet access is needed for this first installation.')
code('python3.12 -m venv .venv-pytorch\nsource .venv-pytorch/bin/activate\npython -m pip install -r requirements.txt')
p('3. Launch the app, then open http://127.0.0.1:8501.')
code('sh start.sh')
p('If python3.12 is not found, install Python 3.12 from python.org, reopen Terminal and repeat the commands. On later launches, use the existing environment and run sh start.sh from the project folder.')
h('On Windows')
p('Open Command Prompt in the extracted project folder. Replace the example folder path with your actual location:')
code('cd /d C:\\path\\to\\ML-IDS-PyTorch-Streamlit\npy -3.12 -m venv .venv-pytorch\n.venv-pytorch\\Scripts\\python -m pip install -r requirements.txt\nstart.bat')
p('Open http://127.0.0.1:8501. Use Command Prompt for the commands shown here. The Windows launcher is included; end-to-end verification was performed on macOS.')
h('Environment and startup checks')
p('From the project folder on macOS, verify the installed application packages:')
code('.venv-pytorch/bin/python -c "import torch, streamlit, shap, scapy; print(\'Ready\')"')
p('If the launcher is not executable, sh start.sh works without changing file permissions. An equivalent direct launch on macOS is:')
code('.venv-pytorch/bin/python -m streamlit run streamlit_app/app.py')
p('For Windows, replace .venv-pytorch/bin/python with .venv-pytorch\\Scripts\\python. Run all commands from the project root so .streamlit/config.toml supplies the light theme, local address and upload limit.')
p('The active project does not require Docker, Kafka, TensorFlow or a live capture driver to open the dashboard and inspect saved packet data.')

page('Prepare raw packet data')
p('The existing /Users/harsha/Downloads/CICIOT23 folder contains train/train.csv, validation/validation.csv and test/test.csv. These files store flow statistics. They cannot be converted back into the original 1,024-byte payload inputs.')
h('1 Obtain the raw CICIoT2023 captures')
p('Use the official dataset page and obtain the raw PCAP edition. The raw download was observed to require registration. Retain the capture labels and source provenance.')
code('https://www.unb.ca/cic/datasets/iotdataset-2023.html')
p('Use verified CICIoT2023 captures. A manifest row must refer to a capture with one verified label. Mixed-label captures require packet-level ground truth and separation before preparation. Keep files from the same original session under the same capture ID.')
h('2 Create and edit the capture manifest')
p('An empty data/packet_manifest.csv is included. If it is missing, choose Create blank manifest in the dashboard or run:')
code('.venv-pytorch/bin/python scripts/packet_pipeline.py init-manifest')
p('Add verified capture rows using docs/packet_manifest.example.csv as a guide, then save as CSV. Example filenames are placeholders. Relative paths start from the manifest folder; absolute paths are accepted. Prepare dataset stays disabled until the manifest and capture paths pass validation.')
table(['Column', 'What to enter'], [
    ('path', 'Existing .pcap or .pcapng file path.'),
    ('label', 'Verified class name, such as BenignTraffic or DDoS-UDP_Flood.'),
    ('split', 'Exactly train, validation or test.'),
    ('capture_id', 'Source session ID. All files from one session stay in one split.'),
    ('dataset', 'Exactly CICIoT2023.'),
    ('source_url', 'Optional provenance URL.')
], [1.1, 5.7])
p('Provide independent capture sessions for all three splits. Every intended class must have usable, distinct payloads in each split. The default normal-traffic label is BenignTraffic. The model can use two or more verified classes; the example template demonstrates two.')
h('3 Prepare the 1024 byte arrays')
p('Activate the project environment, then run:')
code('source .venv-pytorch/bin/activate\npython scripts/packet_pipeline.py prepare --manifest data/packet_manifest.csv')
p('The default output is data/packet_processed. The process scans each capture, reservoir-samples up to 10,000 payload candidates, pads or truncates to 1,024 bytes and preserves capture splits. Duplicate model-visible payloads are excluded from later splits. Empty payloads and IP fragments are skipped.')
p('Expected files: train.npz, validation.npz, test.npz and metadata.json. Existing prepared output is protected. To prepare a different run, add --output data/packet_processed_v2 and use that directory when training.')

page('Train and use the packet model')
h('1 Train and evaluate')
p('After preparation succeeds, run from the project folder with .venv-pytorch activated:')
code('python scripts/packet_pipeline.py train --epochs 25 --batch-size 128')
p('This reads data/packet_processed and writes models/pytorch_packet_ids. The training loop uses PyTorch, seed 42, AdamW and validation-loss checkpoint selection, with early stopping. The held-out test split is evaluated after the best checkpoint is selected. A run may finish before 25 epochs; its duration depends on the data and CPU.')
p('For a different prepared directory, add --data data/packet_processed_v2. If a model already exists, choose a new --output directory to preserve it. In the dashboard, set Model directory to that new directory.')
p('Dashboard alternative: Training & evaluation provides the same preparation and training actions. Set the manifest path, prepared-data directory and new model-output directory; click Prepare dataset, then Start training after preparation finishes. Follow Job log and Learning curves.')
h('2 Confirm the saved output')
table(['Artifact', 'Purpose'], [
    ('model.pt and metadata.json', 'PyTorch weights, model contract, class labels and checksums.'),
    ('status.json and history.json', 'Run state and epoch-by-epoch learning history.'),
    ('evaluation.json', 'Held-out accuracy, precision, recall, F1 and confusion matrices.'),
    ('shap_background.npy', 'Training-partition reference payloads for SHAP.'),
    ('replay.npz', 'Up to 500 held-out payloads for local simulation.')
], [2.35, 4.45])
p('A successful status.json ends with state complete. Refresh the dashboard after training, or restart it with ./start.sh. The sidebar should show “Trained packet model ready”.')
h('3 Run detection and explanations')
bullet('Overview: click Run 5 packets or Start simulation. The app replays held-out payloads through the saved model. Pause simulation to stop replay; export event JSON to save the feed.')
bullet('Packet inspector: choose Held-out sample, PCAP file or Hex payload. Inspect the bytes, then click Classify packet. Uploaded captures are limited to 32 MB and the selector lists up to 200 eligible payloads.')
bullet('Byte explanations: use the selected payload and click Explain packet with SHAP. Inspect signed byte contributions and the sampling residual. Export the explanation as JSON.')
bullet('Training & evaluation: inspect the actual held-out metrics, classwise results and confusion matrix, then export the evaluation JSON.')
h('4 Verify or run from the command line')
code('python -m pytest packet_tests -q\npython scripts/packet_pipeline.py replay --count 10\npython scripts/packet_pipeline.py predict --hex "00 7f ff"\npython scripts/packet_pipeline.py explain --hex "00 7f ff" --samples 256')
p('The test suite uses temporary synthetic fixtures to check software behavior; its results are not CICIoT2023 accuracy results. The three inference commands require real training output. The short example bytes demonstrate command syntax, not a known attack.')

page('Live monitoring and nested validation')
p('The monitoring feed now includes Time, Source IP, Protocol, True classification, Primary model prediction, Confidence, Severity and Status. The interface retains the blue light theme.')
h('1 Observe or analyze packet traffic')
p('Open Live capture & PCAP. Select Live interface, choose your IoT or test-network interface, set a duration and packet limit, then click Start capture. Stop capture ends the run early. The maximum duration is 120 seconds; observed, processed, queued and dropped-packet counts are visible.')
p('Enable Save captured packets locally as PCAP only when you want a local recording. The generated file path appears after capture ends. New live captures are unlabeled and cannot replace the labeled CICIoT2023 training data.')
p('On this Mac, the OS denied access to /dev/bpf0. The UI reports that error. Use a saved authorized capture in the PCAP file tab, or run only the bounded capture CLI with the required OS privileges. Do not run Streamlit as root.')
code('sudo .venv-pytorch/bin/python scripts/packet_pipeline.py capture' + chr(32) + chr(92) + chr(10) + '  --interface en0 --seconds 30 --observe-only' + chr(32) + chr(92) + chr(10) + '  --save-directory runtime/captures')
p('Replace en0 with your actual interface. The OS may request your administrator password. Observe-only mode records metadata without model predictions. Remove --observe-only after real training to classify captured payloads.')
h('2 Read and investigate the feed')
p('For saved data, choose PCAP file, upload a capture, set the analysis limit and click Analyze capture. Use the traffic/status filters or search an address or class. Export CSV or JSON, or select a feed packet and click Use selected packet before opening Packet inspector or Byte explanations.')
p('Source addresses and transport protocols come from packet headers. Port 443 alone does not establish HTTP/TLS. Known held-out labels support MATCH or MISMATCH; live and unlabeled captures show UNVERIFIED. Missing models or payloads show UNCLASSIFIED. Confidence is the actual model probability.')
p('Severity is a documented triage rule: low confidence means Review; benign is Info; DDoS/Mirai is Critical; recon is Medium; other attacks are High. It is not a calibrated risk measurement.')
h('3 Run nested group validation')
p('After data preparation, open Training & evaluation and use Nested validation and bootstrap intervals, or activate .venv-pytorch and run:')
code('python scripts/packet_pipeline.py validate' + chr(32) + chr(92) + chr(10) + '  --outer-folds 5 --inner-folds 3 --bootstrap 100 --epochs 5')
p('The train and validation partitions form the development pool; the deployment test split stays reserved. Inner folds select between two architectures using mean macro F1. Outer folds evaluate the selection procedure. Fixed epochs are used in every fit. The default settings require 35 model fits and enough independent capture groups for every class in every fold.')
p('Outputs are written to output/pytorch/nested_validation. The dashboard displays fold metrics and bootstrap mean, standard deviation and 95 percent percentile bounds. Each bootstrap draw resamples whole capture groups from the out-of-fold predictions; it does not refit a new model.')
p('Real CICIoT2023 results are still pending labeled PCAPs. Nested-validation execution has been tested on temporary synthetic fixtures. It does not create or replace the deployed model checkpoint.')

page('Image review and troubleshooting')
p('The supplied System Architecture and Methodology images describe the intended byte-based design. Use the alignment below when presenting the current application.')
table(['Image component', 'Current application'], [
    ('Scapy ingestion and header stripping', 'Implemented with saved-PCAP analysis, bounded live capture, an inference queue and optional recording. This Mac denied BPF access, so successful live capture still requires OS privileges.'),
    ('1024 byte signal and scaling', 'Implemented. Preserve byte order, keep the first 1,024 bytes, zero-pad shorter payloads and divide byte values by 255.'),
    ('PyTorch CNN BiLSTM', 'Implemented. The BiLSTM processes ordered byte regions within one packet. Real CICIoT2023 training still requires labeled PCAPs.'),
    ('Byte SHAP and Plotly dashboard', 'Implemented and tested with fixtures. Real explanations and detection metrics require the trained packet checkpoint.'),
    ('Nested validation and bootstrap', 'Implemented as a separate experiment with outer 5-fold evaluation, inner 3-fold tuning and 100 capture-level bootstrap draws. Real results require sufficient independent labeled captures.')
], [2.0, 4.8])
p('Presentation corrections: label normalized model inputs as 1,024 float values in [0, 1]; values such as 23 and 255 are pre-normalization byte integers. Scaling is value / 255, not a scaler fitted to dataset minimum and maximum. The metrics and alerts pictured in the slides are illustrative and must not be presented as measured project results.', 'Presentation corrections:')
p('SHAP interpretation: expected gradients estimate contributions to a model probability. They do not establish which bytes actually caused an attack. Empty-payload traffic cannot be scored, and the model does not inspect headers, timing, session context or bytes beyond offset 1023.', 'SHAP interpretation:')
h('Common problems')
p('Port 8501 is already in use. First try opening the existing dashboard. For a separate local instance on macOS, run IDS_PORT=8502 sh start.sh and open http://127.0.0.1:8502.', 'Port 8501 is already in use.')
p('A Python package is missing. Install requirements.txt using the Python executable inside .venv-pytorch. Avoid the historical .venv-macos environment for this application.', 'A Python package is missing.')
p('The model awaits training. Check for models/pytorch_packet_ids/metadata.json and model.pt. Prepare verified raw PCAPs and complete training. Flow CSVs and older TensorFlow checkpoints cannot satisfy this model contract.', 'The model awaits training.')
p('The manifest is missing or empty. Use Create blank manifest or init-manifest if missing, then add verified PCAP rows. Creating the header does not supply training data. The CLI reports an actionable setup error without a Python traceback.', 'The manifest is missing or empty.')
p('Preparation rejects a capture or split. Check paths, labels, capture IDs and split names. Keep sessions separate and provide every required class in each split. Add independent captures if duplicate removal leaves a class empty.', 'Preparation rejects a capture or split.')
p('Output already exists or a checksum fails. Use a fresh output directory for a new run. Keep each model together with its matching metadata, replay and SHAP files; do not combine artifacts from different runs.', 'Output already exists or a checksum fails.')
p('More detail: README.md covers the active project. output/pytorch/dataset_audit.json records the data audit; output/pytorch/delivery_status.json records readiness. runtime/packet-training.log contains jobs launched from the dashboard. The earlier flow dashboard on port 8765 is a separate historical application.', 'More detail:')

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
